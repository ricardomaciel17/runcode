import random
import time
import unicodedata
import requests
import pandas as pd

from bs4 import BeautifulSoup
from concurrent.futures import ThreadPoolExecutor, as_completed
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


def normalize_text(value):
    if value is None:
        return ""

    text = str(value)
    text = text.replace("\xa0", " ")
    text = unicodedata.normalize("NFD", text)
    text = text.encode("ascii", "ignore").decode("utf-8")
    text = " ".join(text.split())

    return text.strip()


def slugify(name):
    return normalize_text(name).lower().replace(" ", "-")


session = requests.Session()
session.headers.update({"User-Agent": "Mozilla/5.0"})


def fetch_last_result(name):

    url = f"https://openresults.run/resultados/{slugify(name)}/"
    print(f"Fetching: {name}")

    try:
        response = session.get(url, timeout=10)

        if response.status_code != 200:
            return None

        soup = BeautifulSoup(response.text, "html.parser")

        event = soup.find("h3", class_="h6")
        if not event:
            return None

        modality = None
        for a in event.find_all("a", href=True):
            href = a["href"]
            if "?modalidade=" in href and "categoria=" not in href:
                modality = normalize_text(a.text)
                break

        table = event.find_next("table")
        if not table:
            return None

        rows = table.find_all("tr")
        if len(rows) < 2:
            return None

        data = [
            normalize_text(td.text)
            for td in rows[1].find_all("td")
        ]

        return {
            "nome": normalize_text(name),
            "modalidade": modality,
            "posicao": data[1],
            "pace": data[4],
            "tempo": data[5],
        }

    except Exception:
        return None


def load_names(filepath):

    df_raw = pd.read_excel(filepath, header=None)

    header_row = next(
        i for i, row in df_raw.iterrows()
        if row.astype(str).str.contains("Nome", case=False, na=False).any()
    )

    df = pd.read_excel(filepath, header=header_row)
    df.columns = df.columns.str.strip()

    return (
        df["Nome"]
        .dropna()
        .apply(normalize_text)
        .tolist()
    )

def worker(name):
    time.sleep(random.uniform(1.4, 3.3) + random.random() * 0.3)
    return fetch_last_result(name)


def process_names(names, max_workers=3):

    results = []
    total = len(names)
    failed = 0

    with ThreadPoolExecutor(max_workers=max_workers) as executor:

        futures = {}

        for i, name in enumerate(names, 1):

            # pausa longa a cada 20 envios
            if i % 20 == 0:
                time.sleep(random.uniform(15, 32))

            future = executor.submit(worker, name)
            futures[future] = name

        for future in as_completed(futures):
            result = future.result()

            if result:
                print(f"{result['nome']} → {result['modalidade']}")
                results.append(result)
            else:
                failed += 1

    return results, total, failed


def pace_to_seconds(pace):
    try:
        m, s = pace.split(":")
        return int(m) * 60 + int(s)
    except:
        return 9999


def extract_number(value):
    try:
        return int("".join(filter(str.isdigit, str(value))))
    except:
        return 9999


def color_results(filepath):

    wb = load_workbook(filepath)
    ws = wb.active

    colors = {
        "green": PatternFill("solid", fgColor="C6EFCE"),
        "blue": PatternFill("solid", fgColor="BDD7EE"),
    }

    for row in ws.iter_rows(min_row=2):

        pos = extract_number(row[2].value)
        pace = row[3].value

        fill = None

        if pos <= 10:
            fill = colors["blue"]
        elif pace_to_seconds(pace) < 240:
            fill = colors["green"]

        if fill:
            for cell in row:
                cell.fill = fill

    wb.save(filepath)


def main():

    input_file = "corrida02.xlsx"
    output_file = "resultados.xlsx"

    names = load_names(input_file)

    results, total, failed = process_names(names)

    pd.DataFrame(results).to_excel(output_file, index=False)

    color_results(output_file)

    print("Finished ✅")
    print(f"Nomes lidos: {total}")
    print(f"Perdas: {failed}")

if __name__ == "__main__":
    main()