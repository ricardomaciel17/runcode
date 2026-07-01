import random
import time
import requests
import xml.etree.ElementTree as ET
import pandas as pd
from concurrent.futures import ThreadPoolExecutor, as_completed
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from bs4 import BeautifulSoup

from utils import normalize_text, slugify, pace_to_seconds, extract_number

session = requests.Session()
session.headers.update({"User-Agent": "Mozilla/5.0"})

def parse_rsf_event(url):
    try:
        r = session.get(url, timeout=10)
        r.encoding = "utf-8"
        root = ET.fromstring(r.text)

        athletes = []
        percursos_set = set()

        for athlete in root.findall(".//Engages/E"):
            nome = athlete.attrib.get("n")
            percurso = athlete.attrib.get("p")
            if nome and percurso:
                athletes.append({"nome": nome, "percurso": percurso})
                percursos_set.add(percurso)

        percursos = sorted(list(percursos_set))
        return athletes, percursos
    except Exception as e:
        print(f"Erro ao processar Clax: {e}")
        return [], []

def fetch_last_result(name):
    url = f"https://openresults.run/resultados/{slugify(name)}/"
    try:
        response = session.get(url, timeout=10)
        if response.status_code != 200:
            return None
        soup = BeautifulSoup(response.text, "html.parser")
        event = soup.find("h3", class_="or-event-card-title")
        if not event:
            event = soup.find("h3", class_="h6")

        if not event:
            return None

        modality = None
        container = event.find_parent()
        if container:
            for a in container.find_all("a", href=True):
                href = a["href"]
                if "?modalidade=" in href and "categoria=" not in href:
                    modality = normalize_text(a.get_text())
                    break

        if not modality:
            for a in soup.find_all("a", href=True):
                href = a["href"]
                if "?modalidade=" in href and "categoria=" not in href:
                    modality = normalize_text(a.get_text())
                    break

        pace = None
        tempo = None
        posicao = None

        stats_container = None
        if container:
            stats_container = container.find(class_="or-athlete-result-stats")
        if not stats_container:
            stats_container = soup.find(class_="or-athlete-result-stats")

        if stats_container:
            for stat in stats_container.find_all(class_="or-athlete-result-stat"):
                label_tag = stat.find("small")
                value_tag = stat.find("strong")
                label = normalize_text(label_tag.get_text()) if label_tag else ""
                value = normalize_text(value_tag.get_text()) if value_tag else ""

                lower_label = label.lower()
                if "pace" in lower_label:
                    pace = value
                elif "líquido" in lower_label or "liquido" in lower_label:
                    tempo = value
                elif "geral" in lower_label:
                    posicao = value
                elif "categoria" in lower_label and not posicao:
                    posicao = value

        if not any([pace, tempo, posicao]):
            table = event.find_next("table")
            if table:
                rows = table.find_all("tr")
                if len(rows) >= 2:
                    data = [normalize_text(td.text) for td in rows[1].find_all("td")]
                    posicao = data[1] if len(data) > 1 else posicao
                    pace = data[4] if len(data) > 4 else pace
                    tempo = data[5] if len(data) > 5 else tempo

        return {
            "nome": normalize_text(name),
            "modalidade": modality,
            "posicao": posicao,
            "pace": pace,
            "tempo": tempo,
        }
    except:
        return None

def worker(name):
    time.sleep(random.uniform(1.2 , 3.3))
    return fetch_last_result(name)

def process_names(names, max_workers=3):
    results = []
    total = len(names)
    failed = 0

    print(f"Iniciando fetching dos resultados para {total} atletas…")

    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = {}

        for i, name in enumerate(names, 1):
            if i % 20 == 0:
                sleep_time = random.uniform(12, 31)
                time.sleep(sleep_time)
            
            print(f"Enviando requisição [{i}/{total}] para {name}…")
            future = executor.submit(worker, name)
            futures[future] = name

        for future in as_completed(futures):
            result = future.result()
            if result:
                results.append(result)
            else:
                failed += 1

    return results, total, failed

def sort_by_priority(results):
    def get_color_priority(result):
        pos = extract_number(result.get("posicao", 9999))
        pace = pace_to_seconds(result.get("pace", "99:99"))

        if pos <= 10:
            return 0
        if pace < 240:
            return 1
        return 2

    painted = [r for r in results if get_color_priority(r) < 2]
    not_painted = [r for r in results if get_color_priority(r) == 2]

    painted_sorted = sorted(
        painted,
        key=lambda r: (get_color_priority(r), normalize_text(r.get("nome", "")).lower()),
    )

    return painted_sorted + not_painted

def color_results(filepath):
    wb = load_workbook(filepath)
    ws = wb.active

    colors = {
        "green": PatternFill("solid", fgColor="C6EFCE"),
        "blue": PatternFill("solid", fgColor="BDD7EE"),
    }

    for row in ws.iter_rows(min_row=2):
        if len(row) < 3:
            continue
        pos = extract_number(row[2].value)
        pace = row[3].value if len(row) > 3 else None
        fill = None

        if pos <= 10:
            fill = colors["blue"]
        elif pace and pace_to_seconds(pace) < 240:
            fill = colors["green"]

        if fill:
            for cell in row:
                cell.fill = fill

    wb.save(filepath)

def run_v2():
    url = input("Digite a URL do evento (RSF ou OnSports): ").strip()
    if not url:
        print("URL inválida.")
        return []

    is_onsports = "onsports" in url.lower() or "mycrono" in url.lower()
    is_runking = "runking" in url.lower() 
    flow_name = "OnSports / RaceZone" if is_onsports else "RSF"
    print(f"Usando fluxo: {flow_name}")

    if is_onsports:
        try:
            r = session.get(url, timeout=10)
            data = r.json()
        except Exception as e:
            print(f"Erro ao abrir JSON OnSports: {e}")
            return []

        athletes = []
        combinacoes_set = set()
        for atleta in data:
            nome = normalize_text(atleta.get("nm"))
            sexo = atleta.get("g")
            sigla = atleta.get("r")
            if nome and sexo and sigla:
                athletes.append({"nome": nome, "sexo": sexo, "sigla": sigla})
                combinacoes_set.add(f"{sexo} {sigla}")

        combinacoes = sorted(list(combinacoes_set))
        if not athletes:
            print("Nenhum atleta encontrado.")
            return []

        print("Grupos disponíveis (sexo + sigla):")
        for i, combo in enumerate(combinacoes, 1):
            print(f"{i} - {combo}")

        choice = input("Escolha o número do grupo (Enter para todos): ").strip()
        if choice.isdigit() and 1 <= int(choice) <= len(combinacoes):
            chosen_combo = combinacoes[int(choice)-1]
            sex_filter, sigla_filter = chosen_combo.split()
        else:
            sex_filter = sigla_filter = None

        filtered_athletes = [
            a for a in athletes
            if (a["sexo"] == sex_filter or not sex_filter) and (a["sigla"] == sigla_filter or not sigla_filter)
        ]

        filtered_names = [a["nome"] for a in filtered_athletes]
        print(f"Atletas filtrados: {len(filtered_names)}")

        results, total, failed = process_names(filtered_names)
    elif is_runking:
        from playwright.sync_api import sync_playwright

        url = "https://resultados.runking.com.br/sportsland/jurere-night-run-hard-rock-cafe-florianopolis-2026?modalityStart=5K&genderStart=M&page=1&pageStart=1"

        with sync_playwright() as p:
            browser = p.chromium.launch(headless=False)
            page = browser.new_page()

            page.goto(url)
            page.wait_for_timeout(5000)

            text = page.inner_text("body") 
            print(text[:2000]) 

            browser.close()
    else:
        athletes, percursos = parse_rsf_event(url)
        if not athletes:
            print("Nenhum atleta encontrado.")
            return []

        print("Percursos disponíveis:")
        for i, p in enumerate(percursos, 1):
            print(f"{i} - {p}")

        choice = input("Escolha o número do percurso (Enter para todos): ").strip()
        if choice.isdigit() and 1 <= int(choice) <= len(percursos):
            chosen_percurso = percursos[int(choice) - 1]
        else:
            chosen_percurso = None

        filtered_names = [
            a["nome"] for a in athletes if (a["percurso"] == chosen_percurso or not chosen_percurso)
        ]
        print(f"Atletas filtrados: {len(filtered_names)}")

        results, total, failed = process_names(filtered_names)

    results_sorted = sort_by_priority(results)
    output_file = "resultados.xlsx"
    pd.DataFrame(results_sorted).to_excel(output_file, index=False)
    color_results(output_file)

    print("Finished ✅")
    print(f"Nomes processados: {len(filtered_names)}")
    print(f"Falhas: {failed}")

    return results_sorted

if __name__ == "__main__":
    run_v2()