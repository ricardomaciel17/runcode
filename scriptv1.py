import random
import time
import unicodedata
import requests
import pandas as pd

from bs4 import BeautifulSoup
from concurrent.futures import ThreadPoolExecutor, as_completed
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


def normalize_text(value):
    if value is None:
        return ""

    text = str(value)
    text = text.replace("\xa0", " ")
    text = unicodedata.normalize("NFD", text)
    text = text.encode("ascii", "ignore").decode("utf-8")
    text = " ".join(text.split())

    return text.strip()


def slugify(name):
    return normalize_text(name).lower().replace(" ", "-")


session = requests.Session()
session.headers.update({"User-Agent": "Mozilla/5.0"})


def fetch_last_result(name):

    url = f"https://openresults.run/resultados/{slugify(name)}/"
    print(f"Fetching: {name}")

    try:
        response = session.get(url, timeout=10)

        if response.status_code != 200:
            return None

        soup = BeautifulSoup(response.text, "html.parser")
        event = soup.find("h3", class_="or-event-card-title")
        if not event:
            event = soup.find("h3", class_="h6")

        if not event:
            return None

        modality = None
        container = event.find_parent()
        if container:
            for a in container.find_all("a", href=True):
                href = a["href"]
                if "?modalidade=" in href and "categoria=" not in href:
                    modality = normalize_text(a.get_text())
                    break

        if not modality:
            for a in soup.find_all("a", href=True):
                href = a["href"]
                if "?modalidade=" in href and "categoria=" not in href:
                    modality = normalize_text(a.get_text())
                    break

        pace = None
        tempo = None
        posicao = None

        stats_container = None
        if container:
            stats_container = container.find(class_="or-athlete-result-stats")
        if not stats_container:
            stats_container = soup.find(class_="or-athlete-result-stats")

        if stats_container:
            for stat in stats_container.find_all(class_="or-athlete-result-stat"):
                label_tag = stat.find("small")
                value_tag = stat.find("strong")
                label = normalize_text(label_tag.get_text()) if label_tag else ""
                value = normalize_text(value_tag.get_text()) if value_tag else ""

                lower_label = label.lower()
                if "pace" in lower_label:
                    pace = value
                elif "líquido" in lower_label or "liquido" in lower_label:
                    tempo = value
                elif "geral" in lower_label:
                    posicao = value
                elif "categoria" in lower_label and not posicao:
                    posicao = value

        if not any([pace, tempo, posicao]):
            table = event.find_next("table")
            if table:
                rows = table.find_all("tr")
                if len(rows) >= 2:
                    data = [normalize_text(td.text) for td in rows[1].find_all("td")]
                    posicao = data[1] if len(data) > 1 else posicao
                    pace = data[4] if len(data) > 4 else pace
                    tempo = data[5] if len(data) > 5 else tempo

        return {
            "nome": normalize_text(name),
            "modalidade": modality,
            "posicao": posicao,
            "pace": pace,
            "tempo": tempo,
        }

    except Exception:
        return None


def load_names(filepath):

    df_raw = pd.read_excel(filepath, header=None)

    header_row = next(
        i for i, row in df_raw.iterrows()
        if row.astype(str).str.contains("Nome", case=False, na=False).any()
    )

    df = pd.read_excel(filepath, header=header_row)
    df.columns = df.columns.str.strip()

    return (
        df["Nome"]
        .dropna()
        .apply(normalize_text)
        .tolist()
    )


def worker(name):
    time.sleep(random.uniform(1.1, 3.6) + random.random() * 0.3)
    return fetch_last_result(name)


def process_names(names, max_workers=3):

    results = []
    total = len(names)
    failed = 0

    with ThreadPoolExecutor(max_workers=max_workers) as executor:

        futures = {}

        for i, name in enumerate(names, 1):

            if i % 20 == 0:
                time.sleep(random.uniform(15, 34))

            future = executor.submit(worker, name)
            futures[future] = name

        for future in as_completed(futures):
            result = future.result()

            if result:
                print(f"{result['nome']} → {result['modalidade']}")
                results.append(result)
            else:
                failed += 1

    return results, total, failed


def pace_to_seconds(pace):
    try:
        m, s = pace.split(":")
        return int(m) * 60 + int(s)
    except:
        return 9999


def extract_number(value):
    try:
        return int("".join(filter(str.isdigit, str(value))))
    except:
        return 9999


def color_results(filepath):

    wb = load_workbook(filepath)
    ws = wb.active

    colors = {
        "green": PatternFill("solid", fgColor="C6EFCE"),
        "blue": PatternFill("solid", fgColor="BDD7EE"),
    }

    for row in ws.iter_rows(min_row=2):

        pos = extract_number(row[2].value)
        pace = row[3].value

        fill = None

        if pos <= 10:
            fill = colors["blue"]
        elif pace_to_seconds(pace) < 240:
            fill = colors["green"]

        if fill:
            for cell in row:
                cell.fill = fill

    wb.save(filepath)


def run_v1():

    input_file = "corrida02.xlsx"
    output_file = "resultados.xlsx"

    names = load_names(input_file)

    results, total, failed = process_names(names)

    def get_priority(result):
        pos = extract_number(result.get("posicao", 9999))
        pace = pace_to_seconds(result.get("pace", "99:99"))

        if pace < 240:
            return 0  
        elif pos <= 10:
            return 1  
        else:
            return 2  

    results_sorted = sorted(results, key=get_priority)

    pd.DataFrame(results_sorted).to_excel(output_file, index=False)

    color_results(output_file)

    print("Finished ✅")
    print(f"Nomes lidos: {total}")
    print(f"Perdas: {failed}")
